
from docx import Document
from io import BytesIO


def generate_comparison_report(result, filename_a, filename_b):
    """Create a downloadable Word comparison report."""

    document = Document()

    document.add_heading(
        "AI Document Comparison Assistant", level=0
    )
    document.add_paragraph(
        "Informational comparison generated using AI."
    )

    document.add_heading("Document A", level=1)
    document.add_paragraph(f"File name: {filename_a}")
    document.add_paragraph(result["summary_a"])

    document.add_heading("Document B", level=1)
    document.add_paragraph(f"File name: {filename_b}")
    document.add_paragraph(result["summary_b"])

    document.add_heading("Similarities", level=1)
    if result["similarities"]:
        for item in result["similarities"]:
            document.add_paragraph(str(item), style="List Bullet")
    else:
        document.add_paragraph("No similarities identified.")

    document.add_heading("Differences", level=1)
    if result["differences"]:
        for item in result["differences"]:
            document.add_heading(
                str(item.get("topic", "Difference")), level=2
            )
            document.add_paragraph(
                f"Change type: {item.get('change_type', 'Not specified')}"
            )
            document.add_paragraph(
                f"Document A: {item.get('document_a', 'Not found in the provided text')}"
            )
            document.add_paragraph(
                f"Document B: {item.get('document_b', 'Not found in the provided text')}"
            )
            document.add_paragraph(
                f"Evidence A: {item.get('evidence_a', 'Not provided')}"
            )
            document.add_paragraph(
                f"Evidence B: {item.get('evidence_b', 'Not provided')}"
            )
    else:
        document.add_paragraph("No meaningful differences identified.")

    document.add_heading("Key Details Comparison", level=1)
    if result["key_details"]:
        for item in result["key_details"]:
            document.add_heading(
                str(item.get("detail", "Detail")), level=2
            )
            document.add_paragraph(
                f"Document A: {item.get('document_a', 'Not found in the provided text')}"
            )
            document.add_paragraph(
                f"Document B: {item.get('document_b', 'Not found in the provided text')}"
            )
    else:
        document.add_paragraph("No key details were returned.")

    document.add_heading("Overall Conclusion", level=1)
    document.add_paragraph(result["conclusion"])

    document.add_heading("Limitations", level=1)
    limitations = result.get("limitations", [])
    if limitations:
        for item in limitations:
            document.add_paragraph(str(item), style="List Bullet")
    else:
        document.add_paragraph(
            "AI-generated results may contain errors. "
            "Verify important details against the original documents."
        )

    document.add_paragraph(
        "This report is informational and is not a legal or compliance decision."
    )

    output = BytesIO()
    document.save(output)
    output.seek(0)

    return output.getvalue()
